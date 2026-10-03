"""Documents → Proposals and Questions.

The model only reads and quotes: it reports what the document states, with a
verbatim quote and dates as written. Code checks every quote against the
text, turns dates into calendar dates, and queues Proposals for review."""

import functools
import io
import json
from pathlib import Path
import re
import unicodedata
import zipfile
from datetime import date, timedelta

from fastapi import APIRouter, BackgroundTasks, HTTPException, Request, UploadFile

from . import asks, inbox, merge, pages, plan
from .clock import local
from .plan import current_term
from .db import WRITE
from .scheduler import notify

router = APIRouter(prefix="/api")

# ---- text extraction ----------------------------------------------------------


def _drop_running_lines(pages: list[str]) -> str:
    """Join pages, removing running headers and footers (a web page printed to
    PDF repeats its title, URL, print time and "3/14" on every page); they
    otherwise split sentences at page breaks and look like dates."""
    shape = lambda line: re.sub(r"\d+", "#", line.strip().lower())
    counts = {}
    for p in pages:
        for s in {shape(l) for l in p.splitlines() if l.strip()}:
            counts[s] = counts.get(s, 0) + 1
    running = {s for s, n in counts.items() if len(pages) >= 3 and n > len(pages) / 2}
    return "\n".join(l for p in pages for l in p.splitlines() if shape(l) not in running)


def _pymupdf(data):
    import pymupdf
    with pymupdf.open(stream=data, filetype="pdf") as doc:
        return _drop_running_lines([page.get_text() for page in doc])


def _pypdf(data):
    from pypdf import PdfReader
    return _drop_running_lines([page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages])


def _python_docx(data):
    import docx
    d = docx.Document(io.BytesIO(data))
    rows = ["\t".join(c.text for c in r.cells) for t in d.tables for r in t.rows]
    return "\n".join([p.text for p in d.paragraphs] + rows)


def _docx_xml(data):
    xml = zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml").decode("utf-8")
    return re.sub(r"<[^>]+>", "", re.sub(r"</w:p>", "\n", xml))


def _plain(data):
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError as e:
            err = e
    raise err


EXTRACTORS = {
    ".pdf": [("PyMuPDF", _pymupdf), ("pypdf", _pypdf)],
    ".docx": [("python-docx", _python_docx), ("raw XML", _docx_xml)],
    ".txt": [("text", _plain)], ".md": [("text", _plain)],
}


def extract(filename: str, data: bytes) -> str:
    """Text of the document. Tries each parser in turn; if all fail, the error
    lists what each one actually said (never a guessed cause)."""
    ext = filename[filename.rfind("."):].lower() if "." in filename else ""
    if ext not in EXTRACTORS:
        raise ValueError(f"Can't read {ext or 'this'} files yet. Upload a PDF, DOCX or TXT.")
    errors = []
    for name, fn in EXTRACTORS[ext]:
        try:
            text = fn(data)
        except Exception as e:
            errors.append(f"{name}: {type(e).__name__}: {e}")
            continue
        if text.strip():
            return text
        errors.append(f"{name}: no text found (the file may be a scanned image)")
    raise ValueError("Couldn't read the file. " + " | ".join(errors))


# ---- model passes -------------------------------------------------------------

RULES = """You read a university course document for a student and report what it states. Rules:
- Report only what the document states explicitly. Never guess, infer or fill in a date, time, room, team or assignment.
- "quote": copy one short passage (under 200 characters) exactly as it appears in the document, word for word, that states the item. Do not paraphrase or fix typos.
- Dates: report them as written. A calendar date → {"type":"date","month":M,"day":D} (add "year" only if written). With a time → {"type":"datetime",...,"time":"HH:MM"} in 24-hour time. "Week N" (optionally a weekday) → {"type":"week","week":N,"weekday":"MO".."SU"}. A date stated relative to another item ("two days before the final") → {"type":"relative","relative_to":"<that item>","offset_days":-2}. Not stated → {"type":"unknown"}.
- Set "provisional": true when the document calls the schedule provisional, tentative or subject to change."""

COURSE_PROMPT = RULES + """
Task: list the course(s) this document is for: course number as department and number only (e.g. "CS 239"; "COM SCI 269" is written "CS 269"; not a term or section code), instructor's full name (only someone the document calls the instructor or professor), title, and the regular weekly meetings: lecture, discussion, lab, seminar, and office hours, each with "type" (lecture, discussion, lab, seminar, office_hours), days as MO,TU,WE,TH,FR,SA,SU, start and end as HH:MM 24-hour, and location (a room, or the online meeting link given for it). If the instructor or a meeting time is not written, leave it empty; do not guess. Quote the line that names the course."""

ITEMS_PROMPT = RULES + """
Task: list everything in this part of the document the student must attend, submit or do:
- "deadline": something the student has to produce or do by a moment: a report, proposal or other submission, a presentation or demo they give, a registration or sign-up.
- "event": a scheduled session the student only attends, with nothing of theirs due, that is not a regular lecture (exam, tutorial, guest lecture, others' presentations).
- "task": a specific piece of work to do before a moment. List each required reading on its own, titled "Read <paper or chapter>", and quote the line that ties it to its lecture or date (e.g. "[Required — Lecture 5]"). Not optional readings (papers teams may choose to present), and not general expectations that apply every week or to whoever presents ("read the papers before class", "participate in discussion", "bring an annotated copy").
- "project": a multi-part deliverable spanning weeks (a course project, a presentation to prepare).
- "question": a date or time only the student can supply: when something of theirs is due or happens when the document leaves it to them (the day they present, their slot). Ask for the date or time itself ("Which day do you present?"), not for the choice behind it (not their topic, paper, team or option). Write the question to ask them, addressed to "you", as a full sentence ending in "?", in "question".
Also set "question" on any other item whose date depends on the student's choice or assignment.
The quote must contain the date you report. When the date is in a heading or table row above the item, quote from the date to the item and mark the skipped middle with "...", e.g. "Lecture 2: Thursday, October 1 ... P2. ReAct".
- "no_class": a specific date the document says there is no class (holiday, break).
- "choice": one thing that happens on one of several dates, depending on the slot the student signs up for or is assigned (e.g. presentations split over two class days). One item, with each date as an entry in "options" ({"when", "quote"}); not one item per date.
Do not list regular lectures, grading percentages, or policies."""

READINGS_PROMPT = RULES + """
Task: the schedule below lists readings (papers, chapters) under class dates. List every reading in it, one per item, titled exactly as the document writes it (join a title broken across lines; copy it word for word). Not session topics or headings ("Model Architecture: Modern Attention Mechanism", "Agents: Coding Agents"), deliverables, links ("Docs, Code"), or "No class" rows."""

READINGS_SCHEMA = {"type": "object", "properties": {"readings": {"type": "array", "items": {"type": "object", "properties": {
    "title": {"type": "string"}}, "required": ["title"]}}}, "required": ["readings"]}

WHEN = {"type": "object", "properties": {
    "type": {"type": "string", "enum": ["date", "datetime", "week", "relative", "weekday", "in_days", "unknown"]},
    "next_week": {"type": "boolean"}, "days": {"type": "integer"},
    "year": {"type": "integer"}, "month": {"type": "integer"}, "day": {"type": "integer"}, "time": {"type": "string"},
    "week": {"type": "integer"}, "weekday": {"type": "string"},
    "relative_to": {"type": "string"}, "offset_days": {"type": "integer"}}, "required": ["type"]}

COURSE_SCHEMA = {"type": "object", "properties": {"courses": {"type": "array", "items": {"type": "object", "properties": {
    "number": {"type": "string"}, "instructor": {"type": "string"}, "title": {"type": "string"}, "quote": {"type": "string"},
    "meetings": {"type": "array", "items": {"type": "object", "properties": {
        "type": {"type": "string", "enum": ["lecture", "discussion", "lab", "seminar", "office_hours"]},
        "days": {"type": "string"}, "start": {"type": "string"}, "end": {"type": "string"}, "location": {"type": "string"},
        "quote": {"type": "string"}}, "required": ["days", "quote"]}}},
    "required": ["number", "instructor", "quote"]}}}, "required": ["courses"]}

ITEMS_SCHEMA = {"type": "object", "properties": {"items": {"type": "array", "items": {"type": "object", "properties": {
    "kind": {"type": "string", "enum": ["deadline", "event", "task", "project", "question", "no_class", "choice"]},
    "title": {"type": "string"}, "course": {"type": "string"}, "quote": {"type": "string"},
    "when": WHEN, "provisional": {"type": "boolean"}, "question": {"type": "string"},
    "options": {"type": "array", "items": {"type": "object", "properties": {"when": WHEN, "quote": {"type": "string"}},
                                           "required": ["when", "quote"]}}},
    "required": ["kind", "title", "quote", "when"]}}}, "required": ["items"]}

QUESTIONS_PROMPT = RULES + """
Task: a careful personal assistant is turning this course document into the student's plan. List the questions it must ask the student before the plan is complete: dates and times only. For example: the day they present or the slot they signed up for, a due date or time the document leaves out for something they must hand in or attend, when a class or session meets if the document doesn't say. Ask for the date or time itself ("Which day do you present?"), never for the choice behind it: not their topic, paper, team, role or which option they take.
Do not ask about course content, grading, or policies. Address the student as "you", one full sentence ending in "?" per question, each with a quote from the document that makes the question necessary. At most 6 questions, most important first."""

# A question the model drafts is asked only if it asks for a date or time.
SCHEDULE = re.compile(r"\b(when|what time|due|deadline|dates?|days?|times?|schedule[ds]?|meets?|week|slot|sign(ed)? up for)\b", re.I)

QUESTIONS_SCHEMA = {"type": "object", "properties": {"questions": {"type": "array", "items": {"type": "object", "properties": {
    "question": {"type": "string"}, "quote": {"type": "string"}}, "required": ["question", "quote"]}}}, "required": ["questions"]}
QUESTIONS_CHARS = 60000


SECOND_LOOK_PROMPT = RULES + """
Task: these items were found in the document, but their quote or their date could not be checked against it. For each numbered item, copy from the document, word for word, the passage that states it and when. If the item and its date are far apart (a heading and a line under it), quote both and mark the skipped middle with "...". Report the date as written ("in your first week" → {"type":"week","week":1}). If the document doesn't say when, quote the passage and give {"type":"unknown"}. If the item isn't in the document, give an empty quote."""

SECOND_LOOK_SCHEMA = {"type": "object", "properties": {"items": {"type": "array", "items": {"type": "object", "properties": {
    "n": {"type": "integer"}, "quote": {"type": "string"}, "when": WHEN}, "required": ["n", "quote", "when"]}}}, "required": ["items"]}


MERGE_PROMPT = """You help a personal assistant decide what to ask a student about one course document. Below are draft questions, numbered. Many ask the same thing in different words.
- Merge drafts that ask the same thing into one question, worded clearly, addressed to "you", ending in "?".
- Drop drafts that don't ask for a date or time (their topic, team, paper, role or choice of option; course content, grading, whether they already did something, reminders).
- Never drop a draft asking when something is due or which day it happens; merge it only with drafts about the same thing.
- Questions marked "already asked" are waiting for the student's answer from another document of the same course. If a draft asks the same thing, put it in that question's group and keep that question's wording; never ask it twice.
- Keep at most 8 questions, most important first.
Return each kept question with "from": the numbers of every draft (and already-asked question) it covers."""

MERGE_SCHEMA = {"type": "object", "properties": {"merged": {"type": "array", "items": {"type": "object", "properties": {
    "question": {"type": "string"}, "from": {"type": "array", "items": {"type": "integer"}}},
    "required": ["question", "from"]}}}, "required": ["merged"]}


def _consolidate(llm, context, drafts, required=(), existing=()):
    """Draft index → (question, quote) to ask, an existing open question (a
    dict) that already asks it, or None if merged away/dropped. `required`:
    drafts that may be merged but never dropped (when is it due, which day).
    `existing`: questions already open for this course from other documents.
    If the model call fails, exact duplicates are merged and the rest kept."""
    same = {q["text"].lower(): q for q in existing}
    exact = {}
    final = [same.get(q.lower()) or exact.setdefault(q.lower(), (q, quote)) for q, quote in drafts]
    if len(drafts) + len(existing) < 2:
        return final
    n = len(drafts)
    listing = "\n".join([f"{i}. {q}" for i, (q, _) in enumerate(drafts)] +
                        [f"{n + j}. {q['text']} (already asked)" for j, q in enumerate(existing)])
    try:
        merged = json.loads(llm.chat([{"role": "system", "content": MERGE_PROMPT},
                                      {"role": "user", "content": f"{context}\n\nDrafts:\n{listing}"}],
                                     schema=MERGE_SCHEMA, timeout=900))["merged"]
    except Exception:
        return final
    out = [None] * n
    for m in merged[:8]:
        ids = [i for i in m.get("from") or [] if isinstance(i, int) and 0 <= i < n + len(existing)]
        members = [i for i in ids if i < n]
        old = next((existing[i - n] for i in ids if i >= n), None)
        q = capitalize((m.get("question") or "").strip())
        if members and (old or q.endswith("?")):
            for i in members:
                out[i] = out[i] or old or (q, drafts[members[0]][1])
    for i in required:  # dropped by the model, but the answer is needed
        out[i] = out[i] or final[i]
    return out


ROLE = re.compile(r"\b(instructors?|professor|prof|lecturer|taught|teacher|faculty)\b")


def named_as_instructor(name: str, text: str) -> bool:
    """The document names this person as the instructor: a role word within a
    line or so of their surname ("Instructor: Stefano Soatto", "Robin Ding /
    Instructor"), not just anywhere ("Thaddy will give a tutorial")."""
    t = _norm_keep_len(text)
    surname = name.split()[-1].lower()
    return any(ROLE.search(t[max(0, m.start() - 60):m.end() + 60])
               for m in re.finditer(rf"\b{re.escape(surname)}\b", t))


def capitalize(s: str) -> str:
    """"gating test" → "Gating test"; "iOS demo" stays."""
    return s[:1].upper() + s[1:] if s[:1].islower() and not s[1:2].isupper() else s


CHUNK = 8000  # characters per items pass; the whole doc would overflow the answer budget


def _ask_model(llm, prompt, schema, context, text):
    raw = llm.chat([{"role": "system", "content": prompt},
                    {"role": "user", "content": f"{context}\n\nDocument:\n\"\"\"\n{text}\n\"\"\""}],
                   schema=schema, timeout=900, max_tokens=6000)
    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        raise ValueError(f"The model returned something that isn't valid JSON: {raw[:200]!r}")


def _items(llm, context, part):
    """The items in one part. A part with a long list (a reading list) can
    overrun the answer: then read it as two halves."""
    try:
        return _ask_model(llm, ITEMS_PROMPT, ITEMS_SCHEMA, context, part)["items"]
    except ValueError:
        halves = chunks(part, len(part) // 2 + 1)
        if len(part) < 2000 or len(halves) < 2:
            raise
        return [it for h in halves for it in _items(llm, context, h)]


def chunks(text, size=CHUNK):
    out, cur = [], ""
    for para in re.split(r"(?<=\n)", text):
        if cur and len(cur) + len(para) > size:
            out.append(cur)
            cur = ""
        cur += para
    return out + [cur] if cur.strip() else out


# ---- checking and resolving ---------------------------------------------------

def _norm(s: str) -> str:
    s = unicodedata.normalize("NFKC", s)
    s = s.translate(str.maketrans({"‘": "'", "’": "'", "“": '"', "”": '"',
                                   "–": "-", "—": "-", " ": " "}))
    return re.sub(r"\s+", " ", s).strip().lower()


_words = lambda s: re.findall(r"\w+", _norm(s))


def _match_at(q, t, start):
    """End index if words q match t from `start`, allowing a few skipped words
    (models drop table cells), else -1. The span may be at most
    len + max(4, len/2) words, so words gathered from different places fail."""
    if start >= len(t) or t[start] != q[0]:
        return -1
    limit, pos = start + len(q) + max(4, len(q) // 2), start
    for w in q[1:]:
        try:
            pos = t.index(w, pos + 1, limit)
        except ValueError:
            return -1
    return pos + 1


ELLIPSIS_GAP = 30  # words a "..." may skip: a table row's cells, not the next row


@functools.lru_cache(maxsize=8)
def _row_starts(text: str) -> frozenset:
    """Word positions where a schedule row starts (a date heading line)."""
    return frozenset(len(_words(text[:m.start()])) for m in DATE_HEADING.finditer(text))


def quoted(quote: str, text: str, gap: int = ELLIPSIS_GAP) -> bool:
    """The quote's words appear in the document in order and close together.
    Pieces joined by "..." may be up to `gap` words apart, within one schedule
    row: "Mon Nov 2 ... Mid-term project report" can't skip over the "Wed Nov 4"
    row the report is in."""
    t = _words(text)
    parts = [p for p in (_words(x) for x in re.split(r"\.\.\.|…", quote or "")) if p]
    if not parts or sum(map(len, parts)) < 2:
        return False
    rows = _row_starts(text) if len(parts) > 1 else frozenset()

    def rest(pieces, frm):
        if not pieces:
            return True
        return any((end := _match_at(pieces[0], t, s)) >= 0 and rest(pieces[1:], end)
                   for s in range(frm, min(frm + gap + 1, len(t))) if not any(frm <= r < s for r in rows))

    return any((end := _match_at(parts[0], t, s)) >= 0 and rest(parts[1:], end) for s in range(len(t)))


def locate(quote: str, text: str, gap: int = ELLIPSIS_GAP):
    """(start, end) character offsets of the quote in the text, else None."""
    spans = [(m.group(), m.start(), m.end()) for m in re.finditer(r"\w+", _norm_keep_len(text))]
    t = [w for w, _, _ in spans]
    parts = [p for p in (_words(x) for x in re.split(r"\.\.\.|…", quote or "")) if p]
    if not parts:
        return None
    for s in range(len(t)):
        end = _match_at(parts[0], t, s)
        for p in parts[1:]:
            if end < 0:
                break
            end = next((e for s2 in range(end, min(end + gap + 1, len(t))) if (e := _match_at(p, t, s2)) >= 0), -1)
        if end >= 0:
            return spans[s][1], spans[end - 1][2]
    return None


LEADING_DATE = re.compile(
    r"^\W*((?:(?:mon|tue|wed|thu|fri|sat|sun)\w*,?\s+)?(?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)\w*\.?\s+\d{1,2}(?:st|nd|rd|th)?\b,?)\s*(.+)$",
    re.I | re.S)


def row_quote(quote: str, text: str) -> str | None:
    """A schedule row quoted as "<date> <deliverable>", skipping the row's
    other cells without marking the gap ("Wed Oct 21 Proposal one-pager",
    with the day's topic and papers in between). Accepted, as
    "<date> … <deliverable>", when both are in the document within a row's
    reach of each other; returns that quote, else None."""
    m = LEADING_DATE.match(quote or "")
    if not m or len(_words(m.group(2))) < 1:
        return None
    fixed = f"{m.group(1).strip()} … {m.group(2).strip()}"
    at = locate(fixed, text)
    if not at:
        return None
    # the same row: no other date line between the date and the deliverable
    between = text[at[0]:at[1]].split("\n")[1:]
    return None if any(DATE_HEADING.match(line) and not DATE_HEADING.match(line).group("week") for line in between) else fixed


SECTION_GAP = 150  # words a "..." may skip within one dated section (a lecture's entry)


def section_quote(quote: str, text: str) -> str | None:
    """A "..." quote whose pieces are further apart than a table row, but in
    the same section: a lecture heading and a note a paragraph below it. No
    other date heading may come between the pieces."""
    if not re.search(r"\.\.\.|…", quote or ""):
        return None
    at = locate(quote, text, SECTION_GAP)
    if not at:
        return None
    between = text[at[0]:at[1]].split("\n")[1:]
    return None if any((m := DATE_HEADING.match(line)) or LECTURE_HEADING.match(line) for line in between) else quote


def checked_quote(quote: str, text: str) -> str | None:
    """The quote as it will be shown, if the document states it: verbatim, a
    schedule row ("<date> … <deliverable>"), or one section ("<heading> … <note>")."""
    if quoted(quote, text):
        return quote
    return row_quote(quote, text) or section_quote(quote, text)


def _norm_keep_len(s: str) -> str:
    """Lower-case with typographic marks unified, same length as the input,
    so word offsets point into the original text."""
    return s.translate(str.maketrans({"‘": "'", "’": "'", "“": '"', "”": '"',
                                      "–": "-", "—": "-", " ": " "})).lower()


MONTHS = ["january", "february", "march", "april", "may", "june", "july", "august", "september", "october",
          "november", "december"]


def date_in_quote(when: dict, quote: str) -> bool:
    """An explicit date counts only if the quote itself states it."""
    w, q = _words(quote), _norm(quote)
    m, d = when.get("month"), when.get("day")
    if not (m and d and 1 <= m <= 12):
        return False
    month_named = any(x in w for x in (MONTHS[m - 1], MONTHS[m - 1][:3], MONTHS[m - 1][:4]))
    numeric = re.search(rf"\b0?{m}[/.-]0?{d}\b", q) is not None
    ordinal = {f"{d}st", f"{d}nd", f"{d}rd", f"{d}th"} & set(w)  # "December 3rd"
    return numeric or (month_named and (str(d) in w or bool(ordinal)))


def time_in_quote(time: str, quote: str) -> bool:
    """"16:00" is stated by "4:00 p.m.", "4pm", "4 PM" or "16:00" in the quote."""
    m = re.fullmatch(r"(\d{1,2}):(\d{2})", (time or "")[:5])
    if not m:
        return False
    h, mm, q = int(m.group(1)), m.group(2), _norm(quote).replace(".", "")
    h12, ampm = (h % 12 or 12), ("pm" if h >= 12 else "am")
    minutes = rf"(:{mm})?" if mm == "00" else rf":{mm}"
    return bool(re.search(rf"\b{h}:{mm}\b", q) or re.search(rf"\b{h12}{minutes}\s*{ampm}\b", q)
                # the start of a range whose am/pm is written once: "4:00-5:50 p.m."
                or re.search(rf"\b{h12}{minutes}\s*-\s*\d{{1,2}}(:\d{{2}})?\s*{ampm}\b", q)
                or (h == 12 and mm == "00" and "noon" in q))


ORDINALS = ["first", "second", "third", "fourth", "fifth", "sixth", "seventh", "eighth", "ninth", "tenth"]
DAY_NAMES = ["monday", "tuesday", "wednesday", "thursday", "friday", "saturday", "sunday"]
NUMBER_WORDS = ["zero", "one", "two", "three", "four", "five", "six", "seven", "eight", "nine", "ten",
                "eleven", "twelve", "thirteen", "fourteen"]


def _explicit(when, today):
    try:
        year = when.get("year") or today.year
        d = date(year, when["month"], when["day"])
        if not when.get("year") and d < today - timedelta(days=60):
            d = date(year + 1, when["month"], when["day"])
    except (ValueError, KeyError, TypeError):
        return None
    return d


def _number_said(n: int, words) -> bool:
    return str(n) in words or (n < len(NUMBER_WORDS) and NUMBER_WORDS[n] in words) or (n == 1 and ("a" in words or "an" in words))


class When:
    """What code could establish about an item's date from its quote."""
    def __init__(self, value=None, window=None, provisional=False, ask=None):
        self.value, self.window, self.provisional, self.ask = value, window, provisional, ask


_WD = r"(monday|tuesday|wednesday|thursday|friday|saturday|sunday|mon|tues|tue|wed|thurs|thur|thu|fri|sat|sun)"
SAID = [  # how a student says a date in chat, read by code (the model's reading can slip)
    (re.compile(r"\bday after tomorrow\b"), lambda m: {"type": "in_days", "days": 2}),
    (re.compile(r"\b(today|tonight)\b"), lambda m: {"type": "in_days", "days": 0}),
    (re.compile(r"\btomorrow\b"), lambda m: {"type": "in_days", "days": 1}),
    (re.compile(r"\bin (\d+|a|an|one|two|three|four|five|six|seven|eight|nine|ten) (day|week)s?\b"),
     lambda m: {"type": "in_days", "days": (int(m[1]) if m[1].isdigit() else max(1, NUMBER_WORDS.index(m[1])) if m[1] in NUMBER_WORDS else 1)
                * (7 if m[2] == "week" else 1)}),
    (re.compile(rf"\bnext {_WD}\b"), lambda m: {"type": "weekday", "weekday": m[1][:2].upper(), "next_week": True}),
    (re.compile(rf"(?<!every )(?<!next )\b{_WD}\b(?!s)"), lambda m: {"type": "weekday", "weekday": m[1][:2].upper()}),
    (re.compile(r"\b(jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\.? (\d{1,2})(st|nd|rd|th)?\b"),
     lambda m: {"type": "date", "month": next(i for i, n in enumerate(MONTHS, 1) if n.startswith(m[1][:3])), "day": int(m[2])}),
    (re.compile(r"\b(\d{1,2})(st|nd|rd|th)? of (jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\b"),
     lambda m: {"type": "date", "month": next(i for i, n in enumerate(MONTHS, 1) if n.startswith(m[3][:3])), "day": int(m[1])}),
    (re.compile(r"\b(\d{1,2})/(\d{1,2})\b"), lambda m: {"type": "date", "month": int(m[1]), "day": int(m[2])}),
]
_CLOCK = re.compile(r"\b(?:at |by |@ ?)?(\d{1,2})(?::(\d{2}))? ?(am|pm|a\.m\.|p\.m\.)|\b(noon)\b|\bat (\d{1,2}):(\d{2})\b")


def said_when(text: str) -> dict | None:
    """The one date a chat message states, as a `when`, or None if it states
    none or more than one (then the model's reading stands)."""
    t, found, taken = _norm(text), [], []
    for pattern, make in SAID:
        for m in pattern.finditer(t):
            if any(a < m.end() and m.start() < b for a, b in taken):
                continue  # "next friday" isn't also "friday"; "day after tomorrow" isn't "tomorrow"
            if m[0] == "may" or (pattern.pattern.startswith(r"\b(\d{1,2})/") and not (1 <= int(m[1]) <= 12)):
                continue
            taken.append((m.start(), m.end()))
            found.append(make(m))
    unique = {json.dumps(w, sort_keys=True) for w in found}
    return found[0] if len(unique) == 1 else None


def said_time(text: str) -> str | None:
    """"at 3pm", "3:30 p.m.", "noon", "at 15:00" → "HH:MM" when the message states one time."""
    times = set()
    for m in _CLOCK.finditer(_norm(text)):
        if m[4]:
            times.add("12:00")
        elif m[5]:
            times.add(f"{int(m[5]):02d}:{m[6]}")
        else:
            h = int(m[1]) % 12 + (12 if m[3].startswith("p") else 0)
            times.add(f"{h:02d}:{m[2] or '00'}")
    return times.pop() if len(times) == 1 else None


def resolve(it: dict, today: date, term: dict | None, known: dict) -> When:
    """Turn the model's report of a date into a calendar date, trusting only
    what the quote states. `known`: lower-case title → date of items resolved
    from the same document (for "two days before Phase 1")."""
    when, quote = it.get("when") or {}, it["quote"]
    t, words = when.get("type"), _words(quote)
    if t in ("date", "datetime"):
        d = _explicit(when, today)
        # stated as a date, or as a weekday the model turned into this week's date ("by Friday" → Oct 2)
        said = d and (date_in_quote(when, quote) or (
            0 <= (d - today).days < 7 and (DAY_NAMES[d.weekday()] in words or DAY_NAMES[d.weekday()][:3] in words)))
        if not said:
            return When()
        time = when.get("time") if t == "datetime" and time_in_quote(when.get("time"), quote) else None
        return When(d.isoformat() + (f"T{time[:5]}" if time else ""))
    if t == "week" and term and isinstance(when.get("week"), int):
        n = when["week"]
        ordinal = ORDINALS[n - 1] if 0 < n <= len(ORDINALS) else None
        if not (re.search(rf"\bweek\s*{n}\b", _norm(quote)) or (ordinal and re.search(rf"\b{ordinal}\s+week\b", _norm(quote)))):
            return When()
        monday = date.fromisoformat(term["week1"]) + timedelta(weeks=n - 1)
        wd = (when.get("weekday") or "").upper()[:2]
        if wd in ("MO", "TU", "WE", "TH", "FR", "SA", "SU"):
            i = ["MO", "TU", "WE", "TH", "FR", "SA", "SU"].index(wd)
            if DAY_NAMES[i] in words or DAY_NAMES[i][:3] in words:
                return When((monday + timedelta(days=i)).isoformat(), provisional=True)
        friday = monday + timedelta(days=4)
        return When(monday.isoformat(), f"{monday}/{friday}", True, asks.which_day(n, monday, friday, it["title"]))
    if t == "weekday" and (when.get("weekday") or "").upper()[:2] in plan.DAYS:
        # "by Friday" = the coming Friday (today counts); "next Friday" = next week's.
        i = plan.DAYS.index(when["weekday"].upper()[:2])
        if not (DAY_NAMES[i] in words or DAY_NAMES[i][:3] in words):
            return When()
        if when.get("next_week"):
            if "next" not in words:
                return When()
            d = today - timedelta(days=today.weekday()) + timedelta(weeks=1, days=i)
        else:
            d = today + timedelta(days=(i - today.weekday()) % 7)
        return When(d.isoformat())
    if t == "in_days" and isinstance(when.get("days"), int) and 0 <= when["days"] <= 366:
        n = when["days"]
        said = {0: {"today", "tonight"}, 1: {"tomorrow"}}.get(n, set()) & set(words) or \
            (_number_said(n, words) and ("day" in words or "days" in words)) or \
            (n % 7 == 0 and _number_said(n // 7, words) and ("week" in words or "weeks" in words)) or \
            (n == 7 and "week" in words)
        return When((today + timedelta(days=n)).isoformat()) if said else When()
    if t == "relative" and isinstance(when.get("offset_days"), int):
        target = (when.get("relative_to") or "").lower().strip()
        base = known.get(target) or next((v for k, v in known.items() if target and (target in k or k in target)), None)
        off = when["offset_days"]
        stated = off == 0 or str(abs(off)) in words or (abs(off) < len(NUMBER_WORDS) and NUMBER_WORDS[abs(off)] in words) \
            or (abs(off) == 7 and "week" in words)
        if base and stated:
            return When((date.fromisoformat(base[:10]) + timedelta(days=off)).isoformat())
    return When()


# ---- ingest -------------------------------------------------------------------

def ingest(con, llm, clock, source_id: int, filename: str, data: bytes = b"", url: str | None = None):
    """Read an uploaded file, or with `url` a course website (the page and the
    sections it links to), and suggest what's in it."""
    try:
        if url:
            filename, text = pages.read_site(url, lambda pdf: extract("page.pdf", pdf))
            with WRITE:
                con.execute("update sources set title = ? where id = ?", (filename, source_id))
        else:
            text = extract(filename, data)
        with WRITE:
            con.execute("update sources set text = ? where id = ?", (text, source_id))
        prior = _supersede(con, source_id)
        _propose_all(con, llm, clock, source_id, text, prior)
        with WRITE:
            con.execute("update sources set status = 'done' where id = ?", (source_id,))
        _tell_done(con, clock, source_id, filename)
    except Exception as e:
        msg = str(e) if isinstance(e, ValueError) else f"The model call failed: {type(e).__name__}: {e}"
        with WRITE:  # withdraw the half-finished result so a retry starts clean
            con.execute("delete from proposals where source_id = ? and status = 'pending'", (source_id,))
            con.execute("delete from questions where source_id = ? and status = 'open'", (source_id,))
            con.execute("update sources set status = 'failed', error = ? where id = ?", (msg, source_id))
        notify(con, clock, "upload", f"Couldn't read {filename}", msg, "#inbox")


def _tell_done(con, clock, source_id, filename):
    """A notification that the suggestions are ready (reading takes minutes)."""
    n = lambda sql: con.execute(sql, (source_id,)).fetchone()[0]
    found = n("select count(*) from proposals where source_id = ? and status = 'pending'")
    asks = n("select count(*) from questions where source_id = ? and status = 'open'")
    about = con.execute("select about from sources where id = ?", (source_id,)).fetchone()[0] or filename
    if not found and not asks:
        return notify(con, clock, "upload", f"Finished reading {about}", "Nothing new to add.", "#inbox")
    parts = [f"{found} suggestion{'s' if found != 1 else ''}"] if found else []
    parts += [f"{asks} question{'s' if asks != 1 else ''} for you in Chat"] if asks else []
    notify(con, clock, "upload", "Suggestions ready" if found else f"Questions about {about}",
           f"{about}: " + " and ".join(parts) + ".", "#inbox" if found else "#chat")


def _supersede(con, source_id) -> dict:
    """If this upload is a new version of a document, withdraw the old
    version's pending proposals and return what earlier versions put in the
    plan: (kind, lower-case title) → current row."""
    lineage = con.execute("select lineage from sources where id = ?", (source_id,)).fetchone()["lineage"]
    if not lineage:
        return {}
    old = [r["id"] for r in con.execute("select id from sources where (id = ? or lineage = ?) and id != ?",
                                        (lineage, lineage, source_id))]
    marks = ",".join("?" * len(old))
    with WRITE:
        con.execute(f"delete from proposals where status = 'pending' and source_id in ({marks})", old)
        con.execute(f"delete from questions where status = 'open' and source_id in ({marks})", old)
        con.execute(f"update sources set replaced_by = ? where id in ({marks}) and replaced_by is null", (source_id, *old))
    prior = {}
    for p in con.execute(f"select id, ops, applied from proposals where status = 'accepted' and source_id in ({marks})", old):
        for op, id in zip(json.loads(p["ops"]), json.loads(p["applied"])):
            if op["op"] != "create" or op["kind"] in ("courses", "terms"):
                continue
            if (plan.origin(con, op["kind"], id) or {}).get("id") != p["id"]:
                continue  # not that version's item any more
            row = con.execute(f"select * from {op['kind']} where id = ?", (id,)).fetchone()
            if row:
                prior[(op["kind"], row["title"].lower())] = dict(row)
    return prior


_MEETING_WORDS = re.compile(r"\b(lectures?|discussions?|labs?|seminars?|sections?|class(es)?|office\s+hours?)\b", re.I)


def _meeting_type(context: str, given=None) -> str:
    c = context.lower()
    if "office hour" in c:
        return "office_hours"
    for t in ("discussion", "lab", "seminar"):
        if t in c:
            return t
    return given if given in ("discussion", "lab", "seminar") else "lecture"


def _meetings_in_text(text):
    """Weekly meetings stated in the text ("Lecture: MW 2:00pm - 3:50pm"),
    for when the model reports none. A link on the next lines is the place."""
    lines = text.splitlines()
    out = []
    for i, line in enumerate(lines):
        around = "\n".join(lines[max(0, i - 1):i + 1])
        if not _MEETING_WORDS.search(around) or not (wk := weekly(line)):
            continue
        link = next((m.group().rstrip(".,)") for l in lines[i:i + 3] if (m := re.search(r"https?://\S+", l))), None)
        out.append({"type": _meeting_type(around), "days": ",".join(wk[0]), "start": wk[1], "end": wk[2],
                    "quote": line.strip(), "location": link})
    return out


def _good_meetings(c, text):
    """The model's weekly meetings whose quote and start time check out; if
    none do, the ones stated in the text."""
    found = _meetings_in_text(text)
    good = []
    for m in c.get("meetings") or []:
        if not quoted(m.get("quote"), text):
            continue
        days = [d for d in (m.get("days") or "").upper().replace(" ", "").split(",") if d in plan.DAYS]
        if not days and (wk := weekly(m["quote"])):  # "MW" instead of "MO,WE": the quote's own days
            m = {**m, "days": ",".join(wk[0])}
        if not time_in_quote(m.get("start"), m["quote"]):  # no usable time: the text's, same days, else ask
            m = next((f for f in found if f["days"] == (m.get("days") or "").upper().replace(" ", "")), m)
        good.append(m)
    return good or found


def _place_checked(m, text):
    """A meeting's location (room or online link) is kept only if the document
    has it; with none given, a link right after its time is the place. Its type
    (lecture, office hours…) is read from its line and the heading above it."""
    loc = (m.get("location") or "").strip()
    if loc and loc.lower() not in text.lower() and not quoted(loc, text):
        loc = ""
    at = locate(m.get("quote") or "", text)
    if at:
        before = text[:at[0]].splitlines()[-1:]  # the heading line above ("Office hours")
        after = text[at[1]:].splitlines()[:3]
        if not loc:
            loc = next((x.group().rstrip(".,)") for l in after if (x := re.search(r"https?://\S+", l))), "")
        mtype = _meeting_type("\n".join(before + [text[at[0]:at[1]]]), m.get("type"))
    else:
        mtype = m.get("type") or "lecture"
    return {**m, "location": loc or None, "type": mtype}


def _expand_choice(it, group):
    """A slot choice ("presentations Nov 30 or Dec 2") → one item per date, grouped,
    so each is planned like any session and one question picks between them."""
    if it.get("kind") != "choice":
        return [it]
    kind = deadline_or_event("event", it.get("title", ""))
    return [{"kind": kind, "title": it.get("title", ""), "course": it.get("course"), "quote": o.get("quote"),
             "when": o.get("when") or {"type": "unknown"}, "slot": group}
            for o in it.get("options") or []]


_SLOTTY = re.compile(r"\b(report|presentations?|present|demo|talk|defen[cs]e|check-?in|slot)\b", re.I)
_SLOT_WORDS = {"first", "second", "half", "part", "continued", "day", "session", "1", "2"}


def _find_slots(items, resolved, keep):
    """The same presentation or report on two or three dates within two weeks
    ("Final project report" Nov 30 and Dec 2) is one slot choice: the student
    is on one of those days. Marks them as a group, like a "choice" item."""
    groups = {}
    for i in keep:
        it = items[i]
        if it["kind"] in ("event", "deadline") and resolved[i].value and _SLOTTY.search(it["title"]) and not it.get("slot"):
            words = frozenset(_words(it["title"])) - STOP - _SLOT_WORDS
            groups.setdefault((it["kind"], words), []).append(i)
    for (_, words), idx in groups.items():
        days = sorted({resolved[i].value[:10] for i in idx})
        if 2 <= len(days) <= 3 and len(days) == len(idx) and \
                (date.fromisoformat(days[-1]) - date.fromisoformat(days[0])).days <= 14:
            for i in idx:
                items[i]["slot"] = "same:" + " ".join(sorted(words))


LECTURES_PROMPT = RULES + """
Task: how this course's class meetings are run.
- "usual": what most class meetings are: "paper_session" (papers or topics are presented and discussed, usually by students taking turns, as in a seminar with graded presentations) or "concept_lecture" (the instructor teaches the material). Quote the passage that shows it.
- "presentation_days": each class meeting in the schedule given to presenting projects (proposals, project reports, demos, final presentations), with its date as written ("when") and a quote that contains the date. A paper presented and discussed in a regular class is not a presentation day. Not days without class, not deliverables handed in outside class."""

LECTURES_SCHEMA = {"type": "object", "properties": {
    "usual": {"type": "object", "properties": {
        "kind": {"type": "string", "enum": ["paper_session", "concept_lecture"]}, "quote": {"type": "string"}},
        "required": ["kind", "quote"]},
    "presentation_days": {"type": "array", "items": {"type": "object", "properties": {
        "when": WHEN, "quote": {"type": "string"}}, "required": ["when", "quote"]}}},
    "required": ["presentation_days"]}


PROJECT_WORDS = re.compile(r"\b(projects?|proposals?|reports?|demos?|posters?|pitch(es)?|milestones?)\b", re.I)  # not "presentation": a paper is presented too


PRESENTED = re.compile(r"\b(presentations?|presented|presents?|reports?|demos?|posters?|pitch(es)?|first half|second half)\b", re.I)


def _schedule_rows(text, today, term):
    """(day or week window, row text) for each schedule row: from a date heading
    ("Wed Oct 14", "Week 9:") to the next."""
    heads = list(DATE_HEADING.finditer(text))
    out = []
    for m, nxt in zip(heads, heads[1:] + [None]):
        row = text[m.start(): nxt.start() if nxt else min(len(text), m.end() + 600)]
        if m.group("week"):
            when = {"type": "week", "week": int(m.group("week"))}
        else:
            month = next(k for k, n in enumerate(MONTHS, 1) if m.group("month").lower()[:3] == n[:3])
            when = {"type": "date", "month": month, "day": int(m.group("day"))}
        w = resolve({"title": "", "quote": m.group(0), "when": when}, today, term, {})
        day = w.window if w.window else (w.value or "")[:10]
        if day:
            out.append((day, row))
    return out


def _lecture_kinds(con, llm, context, text, items, today, term, source_id):
    """Step 9, Lecture kinds (spec: .scratch/lecture-recording): the course's usual
    kind, read by the model and proven by its quote; each class date's papers,
    from the readings already found (a class with papers is a Paper session);
    and presentation days: a schedule row that names a project, report or
    proposal and has no papers is one (code); a row with papers too (Ding's Nov 9:
    "First half Mid-term project report, Second half" a paper) is one when the
    model says so and the row says what's presented. Nothing per course is hard-coded."""
    rows = []
    try:
        got = _ask_model(llm, LECTURES_PROMPT, LECTURES_SCHEMA, context, text[:QUESTIONS_CHARS])
    except Exception:
        got = {}
    usual = got.get("usual") or {}
    if usual.get("kind") and (q := checked_quote(usual.get("quote"), text)):
        rows.append((None, usual["kind"], None, q))
    papers: dict[str, list[str]] = {}
    for it in items:  # all of them: a past lecture (watched as a replay) still has its papers
        if it.get("_class") and re.match(r"(read|review|skim)\b", it["title"], re.I):
            title = re.sub(r"^(read|review|skim)\s+", "", it["title"], flags=re.I)
            if title not in papers.setdefault(it["_class"], []):
                papers[it["_class"]].append(title)
    has_papers = lambda day: any(p == day or ("/" in day and day[:10] <= p <= day[-10:]) for p in papers)
    said = set()  # days the model names as presentation days
    for d in got.get("presentation_days") or []:
        q = checked_quote(d.get("quote"), text)
        w = q and resolve({"title": "", "quote": q, "when": d.get("when")}, today, term, {})
        if w and (w.value or w.window):
            said.add(w.window or w.value[:10])
    shown = set()
    for day, row in _schedule_rows(text, today, term):
        if day in shown or re.search(r"\bno class\b", row, re.I) or not PROJECT_WORDS.search(row):
            continue
        if has_papers(day) and not (day in said and PRESENTED.search(row)):
            continue  # a paper day where something is only handed in ("Proposal one-pager")
        shown.add(day)
        quote = " ".join(row.split())[:160]
        rows.append((day, "presentation_day", json.dumps(papers.pop(day)) if day in papers else None, quote))
    for day, titles in sorted(papers.items()):
        rows.append((day, "paper_session", json.dumps(titles), None))
    with WRITE:
        con.execute("delete from lecture_kinds where source_id = ?", (source_id,))
        con.executemany("insert into lecture_kinds (source_id, date, kind, papers, quote) values (?, ?, ?, ?, ?)",
                        [(source_id, *r) for r in rows])


READING_WORDS = re.compile(r"\b(readings?|reading list|papers?)\b", re.I)
READING_REACH = 150  # words from a class date to the last reading listed under it


def _listed_readings(llm, context, text, today, term) -> list[dict]:
    """Readings a schedule lists under class dates, as tasks dated by the class
    (the day before is set in _settle). The model names the titles; code finds
    each one in the document and the date heading above it, so a title that
    isn't in the document, or isn't under a date, is not planned."""
    heads = [m for m in DATE_HEADING.finditer(text) if m.group("month")]
    if len(heads) < 2 or not READING_WORDS.search(text):
        return []
    schedule = text[heads[0].start():]
    titles = []
    for part in chunks(schedule):
        try:
            titles += [r["title"] for r in _ask_model(llm, READINGS_PROMPT, READINGS_SCHEMA, context, part)["readings"]]
        except Exception:
            continue
    words = _words(text)
    at = []  # (word index, heading, when) for each date heading
    for m in heads:
        month = next(i for i, name in enumerate(MONTHS, 1) if m.group("month").lower()[:3] == name[:3])
        heading = m.group(0).strip(" \t-•*·")
        at.append((len(_words(text[:m.start()])), heading, {"type": "date", "month": month, "day": int(m.group("day"))}))
    out, seen = [], set()
    for title in titles:
        # "… Attention Architecture Paper": the next cell ("Paper Registration") ran in
        title = re.sub(r"\s+paper$", "", re.sub(r"\s+", " ", (title or "").strip()), flags=re.I)
        q = _words(title)
        if len(q) < 2 or tuple(q) in seen:
            continue
        start = next((s for s in range(len(words)) if _match_at(q, words, s) != -1), None)
        under = [h for h in at if start is not None and h[0] <= start]
        if not under or start - under[-1][0] > READING_REACH:
            continue
        seen.add(tuple(q))
        _, heading, when = under[-1]
        out.append({"kind": "task", "title": capitalize(f"Read {title}"), "quote": f"{heading} … {title}",
                    "when": when, "_listed": True})
    return out


def _ask_readings(con, clock, source_id, made):
    """The schedule's readings are suggested held, with one question: does the
    student read them all (for a seminar where each student presents one, they
    may only read their own)."""
    if not made:
        return
    about = con.execute("select about from sources where id = ?", (source_id,)).fetchone()["about"] or ""
    q = inbox.ask(con, clock, source_id, asks.read_each(about.split(":")[0].split(";")[0].strip(), len(made)), None, "choice",
                  None, [{"label": "Yes", "adds": [p["id"] for p in made]}, {"label": "No", "adds": []}])
    with WRITE:
        for p in made:
            con.execute("update proposals set question_id = ? where id = ? and status = 'pending'", (q["id"], p["id"]))


def ask_left_out(con, clock, source_id, dropped):
    """What's still left out after the second look is asked about, not just
    listed: the student says whether it's theirs and when (read like chat).
    Drafted questions whose quote didn't check out aren't asked again."""
    items = [(d["title"], d.get("reason")) for d in dropped if d.get("kind") != "course" and not d["title"].rstrip().endswith("?")]
    if items:
        quote = next((d.get("quote") for d in dropped if d.get("quote")), None) if len(items) == 1 else None
        inbox.ask(con, clock, source_id, asks.left_out(items), quote, "other")


def _ask_slots(con, clock, source_id, slots):
    """For each slot choice with two or more dates: hold the dates and ask which
    one is the student's, with the dates as buttons. The others are dropped."""
    for group in slots.values():
        if len(group) < 2:
            continue
        group.sort(key=lambda g: g[1])
        labels = [_day_label(when) for _, when, _ in group]
        q = inbox.ask(con, clock, source_id, asks.which_slot(group[0][2], labels), None, "choice", None,
                      [{"label": label, "adds": [p["id"]]} for (p, _, _), label in zip(group, labels)])
        with WRITE:
            for p, _, _ in group:
                con.execute("update proposals set question_id = ? where id = ? and status = 'pending'", (q["id"], p["id"]))
            # the open-ended version of this question ("Which date have you signed up for your
            # final project report presentation?") is now asked with buttons: don't ask it twice
            topic = set(_words(group[0][2])) - STOP - _SLOT_WORDS
            for r in con.execute("select id, text from questions where source_id = ? and status = 'open' and id != ? "
                                 "and coalesce(purpose, 'other') = 'other'", (source_id, q["id"])).fetchall():
                said = set(_words(r["text"]))
                if topic <= said and said & {"slot", "date", "day", "signed", "sign", "assigned", "when"}:
                    con.execute("update questions set status = 'dismissed' where id = ?", (r["id"],))


def _day_label(when: str) -> str:
    d = date.fromisoformat(when[:10])
    return f"{d:%a %b} {d.day}"


# The student's rule: something that has them produce or do something (write a
# report, give a presentation, register) is a deadline; something they only
# attend (a guest lecture, an exam) is an event.
TO_DO = re.compile(r"\b(reports?|proposals?|present(s|ing|ations?)?|submi(t|ts|ssions?)|registration|register|sign[- ]?ups?|due"
                   r"|deliverables?|demos?|write[- ]?ups?|one[- ]pagers?|homeworks?|assignments?|drafts?|slides|essays?|posters?|pitch(es)?)\b", re.I)


def course_title_like(t: str) -> bool:
    """"Large Language Models for Code Intelligence", not "CS 239 investigates the design, … agents."."""
    t = t.strip()
    return bool(t) and not t.endswith(".") and len(t.split()) <= 10


def deadline_or_event(kind: str, title: str) -> str:
    """An event the student has to produce something for is a deadline."""
    return "deadline" if kind == "event" and TO_DO.search(title or "") else kind


def _tidy(it):
    """Normalise one reported item, or None for one that isn't ours to plan."""
    # Regular lectures ("Lecture 10: Thursday, November 5 — …") come with the
    # class-meeting Event; models list them anyway.
    if it["kind"] == "event" and re.match(r"\s*lecture\s*\d+\b", it["title"], re.I):
        return None
    if it["kind"] == "task" and re.search(r"\boptional\b", it.get("quote") or "", re.I):
        return None  # optional readings are for whoever picks them, not tasks
    if re.match(r"(suggested|optional)\b", it["title"], re.I) and it["kind"] in ("event", "deadline", "task"):
        return None  # "Suggested team presentations": ideas to choose from, not the student's plan
    it["title"] = capitalize(it["title"].strip())
    if re.match(r"no class\b", it["title"], re.I):
        it["kind"] = "no_class"  # a day off, not something to do
    it["kind"] = deadline_or_event(it["kind"], it["title"])
    return it


def _settle(items, text, today, term, lectures):
    """Dates for the items, and which to plan: (resolved, keep, undated), where
    undated are tasks with no date or lecture (left out unless a second look finds one)."""
    # Resolve absolute dates first so relative ones ("two days before Phase 1") can use them.
    known, resolved = {}, {}
    for rnd in ("absolute", "relative"):
        for i, it in enumerate(items):
            if ((it.get("when") or {}).get("type") == "relative") == (rnd == "relative"):
                # a date found by a fallback (heading above, lecture table) is kept: running
                # the fallback again would start from the quote it already extended
                resolved[i] = it["_fixed"] if "_fixed" in it else resolve(it, today, term, known)
                if resolved[i].value:
                    known.setdefault(it["title"].lower(), resolved[i].value)
    lecture_days = {d for d, _ in lectures.values()}
    for i, it in enumerate(items):
        if it["kind"] == "task" and resolved[i].value and "_fixed" not in it \
                and (resolved[i].value[:10] in lecture_days or it.get("_listed")) \
                and re.match(r"(read|review|skim)\b", it["title"], re.I):
            # dated by its lecture's row: read it the day before
            it["_class"] = resolved[i].value[:10]  # the class it's for (its Lecture kind lists it)
            day = date.fromisoformat(resolved[i].value[:10]) - timedelta(days=1)
            resolved[i] = it["_fixed"] = When(day.isoformat())
            continue
        if resolved[i].value or (it.get("when") or {}).get("type") == "relative":
            continue
        if it["kind"] == "task":
            # A reading is due the day before the lecture it's for.
            if w := _before_lecture(it, lectures) or _before_lecture(it, lectures, _repair(it, text, today, term, known)):
                resolved[i] = it["_fixed"] = w
        elif it["kind"] in ("deadline", "event"):
            # Only schedule entries (deadlines, sessions): general instructions such
            # as "read the papers" sit under headings they don't belong to.
            if w := _repair(it, text, today, term, known):
                resolved[i] = it["_fixed"] = w

    keep, undated, seen = [], [], []
    # slot dates first, then sessions: of a session and its deadline, the session (time, place) is kept
    for i in sorted(range(len(items)), key=lambda i: ("slot" not in items[i], items[i]["kind"] != "event")):
        it = items[i]
        if _duplicate(it, resolved[i], seen):
            continue
        if it["kind"] == "task" and not resolved[i].value and not it.get("question"):
            undated.append(i)
            continue
        # a task whose day has passed is kept, overdue: the reading for a lecture
        # that has already happened is still to be done (the student's call)
        keep.append(i)
    return resolved, sorted(keep), sorted(undated)


def _second_look(llm, context, text, items) -> dict:
    """n → {"quote", "when"} from the model, for items whose quote or date didn't check out."""
    listing = "\n".join(f"{n}. {it['title']} (reported quote: “{it.get('quote') or ''}”)" for n, it in enumerate(items))
    try:
        got = _ask_model(llm, SECOND_LOOK_PROMPT, SECOND_LOOK_SCHEMA, f"{context}\n\nItems:\n{listing}", text[:QUESTIONS_CHARS])["items"]
    except Exception:
        return {}  # a failed second look leaves them as left out
    return {g["n"]: g for g in got if isinstance(g.get("n"), int) and 0 <= g["n"] < len(items) and (g.get("quote") or "").strip()}


def _propose_all(con, llm, clock, source_id, text, prior=None):
    """Read a syllabus in fixed steps (see .scratch/data-rules/spec.md):
    1 the course, 2 weekly meetings, 3 the schedule table (lecture dates),
    4 readings, 5 deliverables, 6 sessions (and slot choices), 7 questions,
    8 proof: every item's quote is checked; misses get one second look."""
    prior = prior or {}
    matched = set()
    today = local(clock.now()).date()
    context = f"Today is {today:%A, %B %d, %Y}."
    dropped = []

    def keep(x, title):
        if quoted(x.get("quote"), text):
            return True
        if (row := row_quote(x.get("quote"), text)):
            x["quote"] = row  # "Wed Oct 21 Proposal one-pager" → "Wed Oct 21 … Proposal one-pager"
            return True
        dropped.append({"title": title, "quote": x.get("quote"), "reason": "quote not found in the document", "kind": "course"})
        return False

    # Pass A: which course(s). Reuse a known Course (same number and instructor).
    courses, meetings, labels = {}, [], []
    title = source_title(con, source_id)
    for c in _ask_model(llm, COURSE_PROMPT, COURSE_SCHEMA, f"{context}\nFile name: {title}", text[:CHUNK])["courses"]:
        c["number"] = course_number(c.get("number") or "")
        if not re.search(r"[A-Za-z].*\d", c["number"]):
            # Not a course number (e.g. the course title). Canvas pages often
            # state it only in the file name: "26F-COM SCI-269-SEM-3 …".
            c["number"] = number_in(title) or number_in(text[:2000]) or ""
        if not c["number"]:
            continue
        instructor = (c.get("instructor") or "").strip()
        if instructor and not named_as_instructor(instructor, text):
            instructor = c["instructor"] = ""  # e.g. a TA named for one tutorial: ask instead
        if not course_title_like(c.get("title") or ""):
            c["title"] = ""  # a sentence from the description, not the course's name
        name = f"{c['number']} · {instructor or 'instructor not stated'}"
        if not keep(c, name):
            continue
        # How questions name this document in chat: course code, instructor (two
        # courses can share a number) and name, not the file name.
        who = f"{c['number']} · {instructor}" if instructor else c["number"]
        label = f"{who}: {c['title'].strip()}" if (c.get("title") or "").strip() else who
        labels.append(label)
        with WRITE:
            con.execute("update sources set about = ? where id = ?", ("; ".join(labels), source_id))
        if not instructor:
            # Course identity needs the instructor; ask rather than guess.
            q = inbox.ask(con, clock, source_id, asks.who_teaches(label), c["quote"], "instructor")
            data = {"number": c["number"], "instructor": "Not stated", **({"title": c["title"]} if c.get("title") else {})}
            p = inbox.propose(con, clock, source_id, f"Add course {c['number']} (edit in the instructor)",
                              [{"op": "create", "kind": "courses", "data": data}], c["quote"], q["id"])
            p = p.get("existing", p)
            inbox.add_target(con, q["id"], p["id"])
            if p.get("status", "pending") == "pending":
                courses[c["number"]] = f"$p{p['id']}.0"
            for m in _good_meetings(c, text):  # the class times stand without the instructor
                meetings.append((courses.get(c["number"]), c["number"], _place_checked(m, text)))
            continue
        known = con.execute("select id from courses where lower(replace(number, ' ', '')) = lower(replace(?, ' ', '')) "
                            "and lower(instructor) = lower(?)", (c["number"], c["instructor"])).fetchone()
        short = f"{c['number']} · {instructor.split()[-1]}"
        if known:
            courses[c["number"]] = known["id"]
        else:
            data = {k: c[k] for k in ("number", "instructor", "title") if c.get(k)}
            same_number = con.execute("select number, instructor from courses where lower(replace(number, ' ', '')) = "
                                      "lower(replace(?, ' ', ''))", (c["number"],)).fetchone()
            q = same_number and inbox.ask(
                con, clock, source_id, asks.same_course(f"{same_number['number']} · {same_number['instructor']}", name),
                c["quote"], "choice")
            p = inbox.propose(con, clock, source_id, f"Add course {name}",
                              [{"op": "create", "kind": "courses", "data": data}], c["quote"], q["id"] if q else None)
            p = p.get("existing", p)  # already pending from an earlier upload: link to that one
            if q:  # "the same course" drops the new one
                inbox.set_options(con, q["id"], [{"label": "A different course", "adds": [p["id"]]},
                                                 {"label": "The same course", "adds": []}])
            if p.get("status") == "pending":
                courses[c["number"]] = f"$p{p['id']}.0"
        for m in _good_meetings(c, text):
            meetings.append((courses.get(c["number"]), short, _place_checked(m, text)))
    only = next(iter(courses.values())) if len(courses) == 1 else None

    # Steps 4-6: readings, deliverables, sessions, a chunk at a time. A slot
    # choice becomes one item per date, grouped. Quotes that don't check out
    # get a second look (step 8).
    items, missed = [], []
    for part in chunks(text):
        for n, it in enumerate(_items(llm, context, part)):
            for x in _expand_choice(it, f"{len(items) + len(missed)}.{n}"):
                if (q := checked_quote(x.get("quote"), text)):
                    items.append(_tidy({**x, "quote": q}))
                else:
                    missed.append(_tidy(x))
    items = [it for it in items if it]
    missed = [it for it in missed if it]

    # Step 3: the schedule table gives each lecture's date (readings are due the day before).
    term = current_term(con, today)
    lectures = lecture_dates(text, today, term)
    resolved, keep_items, undated = _settle(items, text, today, term, lectures)

    # Step 8, coverage: schedule rows that list something due or happening
    # ("Fri Oct 9 No Class — Project Team List Due") but that nothing read
    # lands on are read again on their own, then checked like everything else.
    if rows := _uncovered_rows(text, items, resolved, keep_items, today, term):
        try:
            got = _ask_model(llm, ROWS_PROMPT, ITEMS_SCHEMA, context, "\n\n".join(rows))["items"]
        except Exception:
            got = []
        for n, it in enumerate(got):
            for x in _expand_choice(it, f"rows.{n}"):
                if (q := checked_quote(x.get("quote"), text)) and (x := _tidy({**x, "quote": q})):
                    items.append(x)
        resolved, keep_items, undated = _settle(items, text, today, term, lectures)

    # Step 4b, the reading list: a schedule that lists papers under class dates
    # ("Readings are listed by lecture date") but that the item pass read as topics
    # (or as papers to present) is read for its readings alone. Code finds the
    # class each title sits under; each reading is due the day before.
    if not any(it["kind"] == "task" and re.match(r"read\b", it["title"], re.I) for it in items):
        if listed := _listed_readings(llm, context, text, today, term):
            items += listed
            resolved, keep_items, undated = _settle(items, text, today, term, lectures)

    # Second look: what was left out (quote not found, or no date) goes back to
    # the model once, for the exact passage and its date. Whatever checks out
    # is planned like everything else; the rest is listed as left out.
    again = [(it, None) for it in missed] + [(items[i], i) for i in undated]
    found = _second_look(llm, context, text, [it for it, _ in again]) if again else {}
    for n, (it, at) in enumerate(again):
        g = found.get(n)
        q = g and checked_quote(g.get("quote"), text)
        if q:
            fixed = {**it, "quote": q, "when": g.get("when") or it.get("when") or {"type": "unknown"}}
            if at is None:
                items.append(fixed)
            else:
                items[at] = fixed
        elif at is None:
            dropped.append({"title": it["title"], "quote": it.get("quote"), "reason": "quote not found in the document"})
    if found:
        resolved, keep_items, undated = _settle(items, text, today, term, lectures)
    for i in undated:
        # No date and no lecture even on a second look: a general instruction, not a task to plan.
        dropped.append({"title": items[i]["title"], "quote": items[i]["quote"], "reason": "no date or lecture stated"})

    _find_slots(items, resolved, keep_items)

    # Sessions stated as a weekly time ("Lecture: MW 2:00pm - 3:50pm") become weekly events below.
    weekly_ids = {i for i in keep_items if items[i]["kind"] == "event" and not resolved[i].value and weekly(items[i]["quote"])}
    weekly_items = []

    # Step 7. Questions: what only the student can tell us, from a pass of its own (so
    # asking doesn't depend on the item pass remembering to), from the items,
    # and the ones code must ask (no date; only a week). All are drafts, merged
    # with each other and with what's already open for this course, then asked.
    drafts, item_draft, week_draft, required = [], {}, {}, []

    def draft(q, quote, check=True):
        q = capitalize((q or "").strip())
        if check and not SCHEDULE.search(q):
            return None  # only dates and times are asked; not their topic, team or choices
        if (q.endswith("?") or not check and "?" in q) and not re.match(r"(have|did) you\b", q, re.I) and (not check or quoted(quote, text)):
            drafts.append((q, quote))
            return len(drafts) - 1

    for q in _ask_model(llm, QUESTIONS_PROMPT, QUESTIONS_SCHEMA, context, text[:QUESTIONS_CHARS])["questions"]:
        if draft(q.get("question"), q.get("quote")) is None and q.get("question") and not quoted(q.get("quote"), text):
            dropped.append({"title": q["question"], "quote": q.get("quote"), "reason": "quote not found in the document"})
    for i in keep_items:
        it, w = items[i], resolved[i]
        kind = KIND.get(it["kind"], (None,))[0]
        d = draft(it.get("question") or (it["title"] if it["kind"] == "question" else ""), it["quote"])
        if d is None and kind in ("deadlines", "events", "projects") and not w.value and i not in weekly_ids:
            # A deliverable or session without a stated date must be asked about.
            d = draft(asks.when_is(it["title"], kind), it["quote"], False)
            required.append(d)
        if d is not None:
            item_draft[i] = d
        if w.ask and kind != "tasks":  # which day of the week: worth asking for a session or deadline; a task just has the week
            week_draft[i] = draft(w.ask, it["quote"], False)
            required.append(week_draft[i])
    about = con.execute("select about from sources where id = ?", (source_id,)).fetchone()["about"]
    existing = [dict(r) for r in con.execute(
        "select q.* from questions q join sources s on s.id = q.source_id where q.status = 'open' "
        "and s.about = ? and s.id != ?", (about, source_id))] if about else []
    final = _consolidate(llm, context, drafts, required, existing)
    asked = {}

    def ask_draft(d):
        if d is None or final[d] is None:
            return None
        if isinstance(final[d], dict):  # already open from another document
            return final[d]
        q, quote = final[d]
        if q not in asked:
            # code's own questions (when is it due, which day) are for a date; the model's are read like chat
            asked[q] = inbox.ask(con, clock, source_id, q, quote, "date" if d in required else "other")
        return asked[q]

    for d in range(len(drafts)):
        ask_draft(d)

    no_class, slots, readings = [], {}, []
    for i in keep_items:
        it = items[i]
        if i in weekly_ids and (wk := weekly(it["quote"])):
            # "Lecture: MW 2:00pm - 3:50pm" is a weekly class, not a one-off with a missing date
            course = courses.get(course_number(it.get("course") or "")) or only
            number = next((n for n, v in courses.items() if v == course), None) or (it.get("course") or "")
            lecture = re.match(r"(lectures?|class(es)?|meetings?)\b", it["title"], re.I)
            m = {"days": ",".join(wk[0]), "start": wk[1], "end": wk[2], "quote": it["quote"]}
            weekly_items.append((course, number, m, None if lecture else f"{number} {it['title'].lower()}".strip()))
            continue
        if it["kind"] == "no_class":
            if resolved[i].value and not resolved[i].window:
                no_class.append(resolved[i].value[:10])
            continue
        course = courses.get(course_number(it.get("course") or "")) or only
        held = ask_draft(item_draft.get(i))
        made = _propose_item(con, clock, source_id, it, course, resolved[i], prior, matched, held)
        for q in (held, ask_draft(week_draft.get(i))):  # their answers act on this item ("which day in Week 9…?")
            if made and q:
                inbox.add_target(con, q["id"], made["id"])
        if made and it.get("slot") and resolved[i].value:
            slots.setdefault(it["slot"], []).append((made, resolved[i].value, it["title"]))
        if made and it.get("_listed"):
            readings.append(made)
    _ask_slots(con, clock, source_id, slots)
    _ask_readings(con, clock, source_id, readings)

    for course, short, m in meetings:
        office = m.get("type") == "office_hours"
        p = _propose_meetings(con, clock, source_id, course, short, m, term, no_class, prior, matched,
                              f"{short.split(' · ')[0]} office hours" if office else None)
        if p and office:
            inbox.mark_optional(con, p["id"])  # offered unticked: the student decides
    for course, short, m, title in weekly_items:
        _propose_meetings(con, clock, source_id, course, short, m, term, no_class, prior, matched, title)

    # Whatever earlier versions added that this version no longer mentions.
    for (kind, _), row in prior.items():
        if (kind, row["title"].lower()) not in matched:
            inbox.propose(con, clock, source_id, f"Remove “{row['title']}”? It's not in the new version of {title}.",
                          [{"op": "delete", "kind": kind, "id": row["id"]}])

    _lecture_kinds(con, llm, context, text, items, today, term, source_id)

    with WRITE:
        con.execute("update sources set dropped = ? where id = ?", (json.dumps(dropped), source_id))
    ask_left_out(con, clock, source_id, dropped)


HEADING_LINES = 15  # how far above an item its date heading may be (one table row)
_MONTH_RE = "|".join(m[:3] + r"\w*" for m in MONTHS)
# A date that starts a line (a table row or heading), optionally after a short
# label such as "Lecture 2:". Dates inside prose sentences don't qualify.
DATE_HEADING = re.compile(
    rf"^[^\w\n]*(?:[A-Za-z]+\s+\d{{1,2}}\s*:\s*)?"
    rf"(?:(?:(?:mon|tue|wed|thu|fri|sat|sun)\w*,?\s+)?(?P<month>{_MONTH_RE})\.?\s+(?P<day>\d{{1,2}})\b(?!\s*,?\s*\d{{4}}\s*,?\s*\d{{1,2}}:)"
    rf"|week\s+(?P<week>\d{{1,2}})\b)", re.I | re.M)


ROWS_PROMPT = RULES + """
Task: each part below is one row of the course schedule, which the first reading found nothing in. These rows were picked because they name something the student submits, takes or presents: a list or report to hand in, a test or exam, a report or presentation day, a check-in, a sign-up. Report each such thing as an item ("deadline" for something the student hands in or presents, "event" for something they only attend or take, such as an exam, "choice", or "no_class"), even when the row doesn't say "due", with a quote from the row that includes its date, marking a skipped middle with "...". For example: "Fri Oct 9 No Class Project Team List" → no_class and a deadline "Project team list"; "Wed Nov 4 Mid-term project report" → a deadline "Mid-term project report"; "Week 1: Introduction, gating test" → an event "Gating test" in week 1. Not lecture topics, readings, or suggested/optional presentations."""

ROW_WORDS = re.compile(r"\b(due|deadline|submit\w*|report|proposal|team list|test|exam|quiz|midterm|final|check-?in|"
                       r"registration|register|sign-?up|assessment|deliverables?|demo)\b", re.I)


def _uncovered_rows(text, items, resolved, keep, today, term) -> list[str]:
    """Dated schedule rows (a date or "Week N" heading and what follows it) that
    mention a deliverable or session, on days no item read so far lands on."""
    covered = set()
    for i in keep:
        w = resolved[i]
        if items[i]["kind"] == "no_class" or not w.value:
            continue
        a, b = (w.window.split("/") if w.window else (w.value[:10], w.value[:10]))
        d = date.fromisoformat(a)
        while d.isoformat() <= b:
            covered.add(d.isoformat())
            d += timedelta(days=1)
    heads = list(DATE_HEADING.finditer(text))
    out = []
    for k, m in enumerate(heads):
        end = heads[k + 1].start() if k + 1 < len(heads) else len(text)
        row = text[m.start():min(end, m.start() + 600)].strip()
        if not ROW_WORDS.search(row):
            continue
        if m.group("week"):
            when = {"type": "week", "week": int(m.group("week"))}
        else:
            month = next(i for i, name in enumerate(MONTHS, 1) if m.group("month").lower()[:3] == name[:3])
            when = {"type": "date", "month": month, "day": int(m.group("day"))}
        w = resolve({"title": "", "quote": m.group(0).strip(), "when": when}, today, term, {})
        if not w.value:
            continue
        first, last = (w.window[:10], w.window[11:]) if w.window else (w.value[:10], w.value[:10])
        if any(first <= x <= last for x in covered):
            continue
        out.append(row)
    return out[:20]


def _repair(it, text, today, term, known):
    """An item whose own quote states no date takes the nearest date heading
    above it (schedules put the date on the row or heading, not beside every
    item). Deterministic on purpose: a model asked to pick a date from the
    passage picked the wrong row's date and a print timestamp. The quote shown
    for review then starts with that heading, so the reason is visible."""
    at = locate(it["quote"], text)
    if not at:
        return None
    above = text[:at[0]].split("\n")
    passage = "\n".join(above[-HEADING_LINES:])
    heading = None
    for m in DATE_HEADING.finditer(passage):
        heading = m
    if not heading:
        return None
    if heading.group("week"):
        when = {"type": "week", "week": int(heading.group("week"))}
    else:
        month = next(i for i, name in enumerate(MONTHS, 1) if heading.group("month").lower()[:3] == name[:3])
        when = {"type": "date", "month": month, "day": int(heading.group("day"))}
    quote = heading.group(0).strip(" \t-•*·")
    w = resolve({"title": it["title"], "quote": quote, "when": when}, today, term, known)
    if not w.value:
        return None
    it["quote"] = f"{quote.strip()} … {it['quote'].strip()}"
    return w


LECTURE_HEADING = re.compile(
    rf"^[^\w\n]*(?:lecture|class|session)\s+(?P<n>\d{{1,2}})\s*[:.\-–—]\s*"
    rf"(?:(?:mon|tue|wed|thu|fri|sat|sun)\w*,?\s+)?(?P<month>{_MONTH_RE})\.?\s+(?P<day>\d{{1,2}})\b", re.I | re.M)
LECTURE_REF = re.compile(r"\b(?:lecture|class|session)\s*[:#\-–—]?\s*(\d{1,2})\b", re.I)


def lecture_dates(text, today, term) -> dict:
    """Lecture number → (date, heading) from schedule lines such as
    "Lecture 5: Tuesday, October 13 — Agent-Computer Interfaces"."""
    out = {}
    for m in LECTURE_HEADING.finditer(text):
        month = next(i for i, name in enumerate(MONTHS, 1) if m.group("month").lower()[:3] == name[:3])
        heading = m.group(0).strip(" \t-•*·")
        w = resolve({"title": "", "quote": heading, "when": {"type": "date", "month": month, "day": int(m.group("day"))}},
                    today, term, {})
        if w.value:
            out.setdefault(int(m.group("n")), (w.value[:10], heading))
    return out


def _before_lecture(it, lectures, under=None) -> When | None:
    """A reading for Lecture N ("[Required — Lecture 5]") is due the day
    before that lecture. `under`: the schedule row the item sits under, for
    a reading listed in the schedule itself (titles starting "Read")."""
    if under:
        if not re.match(r"(read|review|skim)\b", it["title"], re.I):
            return None
        day = date.fromisoformat(under.value[:10])
    else:
        m = LECTURE_REF.search(_norm(it["quote"])) or LECTURE_REF.search(_norm(it["title"]))
        if not m or int(m.group(1)) not in lectures:
            return None
        iso, heading = lectures[int(m.group(1))]
        day = date.fromisoformat(iso)
        it["quote"] = f"{heading} … {it['quote'].strip()}"
    it["_class"] = day.isoformat()
    return When((day - timedelta(days=1)).isoformat())


COMPARED = ("due", "start", "end", "deadline", "window", "repeat", "until", "skip", "location", "provisional")


def _emit(con, clock, source_id, kind, data, quote, question_id, prior, matched, summary=None):
    """Propose `data` as a new item, or, if an earlier version of this
    document already put the same item in the plan, only what changed."""
    key = (kind, data["title"].lower())
    old = prior.get(key)
    if not old:  # new to this document: but maybe not to the plan (Bruin Learn, another file, chat)
        return merge.propose_or_fill(con, clock, source_id, kind, data, quote, question_id, summary)
    matched.add(key)
    changed = {f: data[f] for f in COMPARED if f in data and data[f] != old.get(f)
               and not (f == "provisional" and bool(data[f]) == bool(old.get(f)))}
    if changed:
        what = "; ".join(f"{f} {old.get(f) or 'none'} → {v}" for f, v in changed.items())
        p = inbox.propose(con, clock, source_id, f"Update “{old['title']}”: {what}",
                          [{"op": "update", "kind": kind, "id": old["id"], "data": changed}], quote, question_id)
        return p if "id" in p else None


STOP = {"the", "a", "an", "and", "or", "of", "for", "to", "in", "on", "due", "your", "with"}


def _duplicate(it, when, seen) -> bool:
    """One item reported twice in the same document, by the one duplicate rule
    (merge.same): e.g. "Phase 1 due" and "Phase 1 application, traces and draft
    requirements" on the same day."""
    table = KIND.get(it["kind"], (None,))[0]
    if not table or not when.value:
        return False
    me = {"title": it["title"], DATE_FIELD_OF[table]: when.value, "window": when.window}
    dup = any(merge.same(t, other, table, me) for t, other in seen)
    seen.append((table, me))  # remembered even when dropped, so a third report matches through it
    return dup


_DAY_RUN = re.compile(r"(m|tu|t|w|th|r|f|sa|su)+")


def weekly(words_text: str):
    """(days, start, end) of a weekly time stated in words: "MW 2:00pm - 3:50pm",
    "Tuesdays and Thursdays, 4-5:50 p.m.", "TTh 10am". None if days or a start
    time aren't both there. Times are 24-hour "HH:MM"; end may be None."""
    q = _norm(words_text).replace(".", "")
    days, repeating = [], "every" in q.split()
    for w in re.findall(r"[a-z]+", q):
        for i, name in enumerate(DAY_NAMES):
            if w.startswith(name[:3]) and (w in (name, name + "s", name[:3], name[:4]) or w.startswith(name)):
                days.append(plan.DAYS[i])
                repeating |= w == name + "s"  # "Mondays"
        if len(w) >= 2 and _DAY_RUN.fullmatch(w) and w not in ("am", "pm", "th"):
            repeating = True  # "MW", "TTh"
            for tok in re.findall(r"th|tu|sa|su|[mtwrf]", w):
                days.append({"m": "MO", "t": "TU", "tu": "TU", "w": "WE", "th": "TH", "r": "TH", "f": "FR", "sa": "SA", "su": "SU"}[tok])
    days = [d for d in plan.DAYS if d in days]
    m = re.search(r"\b(\d{1,2})(?::(\d{2}))?\s*(am|pm)?\s*(?:-|to)\s*(\d{1,2})(?::(\d{2}))?\s*(am|pm)\b", q) \
        or re.search(r"\b(\d{1,2})(?::(\d{2}))?\s*(am|pm)\b()()()", q)
    if not days or not m:
        return None
    h1, m1, ap1, h2, m2, ap2 = m.groups()
    if not (repeating or h2):  # "Due Friday at 5pm" is one day, not every week
        return None
    ap1 = ap1 or ap2

    def hhmm(h, mm, ap):
        h = int(h) % 12 + (12 if ap == "pm" else 0)
        return f"{h:02d}:{mm or '00'}"
    return days, hhmm(h1, m1, ap1), (hhmm(h2, m2, ap2) if h2 else None)


def meeting_data(term, days, start, end, title, course=None, location=None, no_class=()) -> dict:
    """A weekly Event over the term's instruction weeks, skipping holidays on those days."""
    first = date.fromisoformat(term["week1"]) - timedelta(days=7)
    while first.isoformat() < term["instruction_begins"] or plan.DAYS[first.weekday()] not in days:
        first += timedelta(days=1)
    holidays = [h[:10] for h in term["holidays"].splitlines()
                if plan.DAYS[date.fromisoformat(h[:10]).weekday()] in days]
    data = {"title": title, "start": f"{first}T{start[:5]}", "repeat": ",".join(days), "until": term["instruction_ends"]}
    if end:
        data["end"] = f"{first}T{end[:5]}"
    if location:
        data["location"] = location
    if course is not None:
        data["course_id"] = course
    if skip := sorted(set(holidays + list(no_class))):
        data["skip"] = ",".join(skip)
    return data


def _propose_meetings(con, clock, source_id, course, short, m, term, no_class, prior, matched, title=None):
    """Regular class meetings → one weekly Event over the term's instruction
    weeks, skipping holidays on those days and the document's no-class days."""
    days = [d for d in (m.get("days") or "").upper().replace(" ", "").split(",") if d in plan.DAYS]
    start = m.get("start") if time_in_quote(m.get("start"), m["quote"]) else None
    end = m.get("end") if time_in_quote(m.get("end"), m["quote"]) else None
    if not days or not term:
        return None
    if not start:
        inbox.ask(con, clock, source_id, asks.meeting_time(short, days), m["quote"], "meeting",
                  {"course": course, "short": short, "days": days, "title": title})
        return None
    data = meeting_data(term, days, start, end, title or f"{short.split(' · ')[0]} class", course, m.get("location"), no_class)
    return _emit(con, clock, source_id, "events", data, m["quote"], None, prior, matched,
                 f"{title} (weekly)" if title else f"{short} class meetings")


DATE_FIELD_OF = {"deadlines": "due", "events": "start", "tasks": "due", "projects": "deadline"}
KIND = {"deadline": ("deadlines", "due"), "event": ("events", "start"), "task": ("tasks", "due"),
        "project": ("projects", "deadline")}


def source_title(con, source_id) -> str:
    return con.execute("select title from sources where id = ?", (source_id,)).fetchone()["title"]


def number_in(s: str) -> str | None:
    """A course number written in text: "COM SCI-269", "CS239", "CS 239" → normalised."""
    m = re.search(r"\b([A-Z]{2,}(?:\s[A-Z]{2,})*)[\s-]?(\d{2,3}[A-Z]?)\b", s)
    return course_number(f"{m.group(1)} {m.group(2)}") if m else None


course_number = plan.course_number  # "COM SCI 269" → "CS 269"


def _propose_item(con, clock, source_id, it, course, when: When, prior, matched, question=None):
    """question: the (merged) question the item waits on, if the model drafted one."""
    kind, field = KIND.get(it["kind"], (None, None))
    if not kind:
        return None
    data = {"title": it["title"]}
    if when.value:
        # a task "in Week 1" is due by the end of that week, not its Monday
        data[field] = when.window[-10:] if kind == "tasks" and when.window else when.value
    if when.window and kind != "projects":
        data["window"] = when.window
    if course is not None:
        data["course_id"] = course
    if it.get("provisional") or when.provisional:
        data["provisional"] = True
    if kind == "events" and "start" not in data:
        # An Event needs a time; without one it's a Deadline-like marker until answered.
        kind = "deadlines"
    return _emit(con, clock, source_id, kind, data, it["quote"], question["id"] if question else None, prior, matched)


# ---- HTTP ---------------------------------------------------------------------

def _doc_name(filename: str) -> str:
    """"CS239 (1).pdf" and "cs239.pdf" name the same document."""
    stem = filename.rsplit(".", 1)[0]
    return re.sub(r"\s*\(\d+\)$", "", stem).strip().lower()


@router.post("/uploads", status_code=202)
async def upload(file: UploadFile, background: BackgroundTasks, request: Request, replaces: int | None = None):
    """Upload a document; `replaces` (or the same file name) makes it a new
    version of an earlier upload."""
    s = request.app.state
    data = await file.read()
    now = local(s.clock.now()).strftime("%Y-%m-%dT%H:%M")
    if replaces is None:
        prev = next((r for r in s.db.execute("select id, title from sources where kind = 'document' and url is null and replaced_by is null "
                                             "and status = 'done' order by id desc") if _doc_name(r["title"]) == _doc_name(file.filename)), None)
    else:
        prev = s.db.execute("select id from sources where id = ?", (replaces,)).fetchone()
        if not prev:
            raise HTTPException(404, "The upload to replace doesn't exist.")
    lineage = None
    if prev:
        lineage = s.db.execute("select coalesce(lineage, id) from sources where id = ?", (prev["id"],)).fetchone()[0]
    with WRITE:
        cur = s.db.execute("insert into sources (kind, title, text, status, created_at, lineage) "
                           "values ('document', ?, '', 'processing', ?, ?)", (file.filename, now, lineage))
    _kept(s, cur.lastrowid, file.filename).write_bytes(data)  # so a failed upload can be retried
    background.add_task(ingest, s.db, s.reader, s.clock, cur.lastrowid, file.filename, data)
    return {"id": cur.lastrowid, "status": "processing"}


@router.post("/uploads/url", status_code=202)
async def upload_url(background: BackgroundTasks, request: Request, replaces: int | None = None):
    """Read a course website from its address; the same address again (or
    `replaces`) is a new version of it."""
    s = request.app.state
    url = ((await request.json()).get("url") or "").strip()
    if not re.match(r"https?://[^\s/]+\.[^\s]+$", url):
        raise HTTPException(422, "That isn't a web address (it should start with http:// or https://).")
    now = local(s.clock.now()).strftime("%Y-%m-%dT%H:%M")
    prev = s.db.execute("select id from sources where id = ?", (replaces,)).fetchone() if replaces is not None else \
        s.db.execute("select id from sources where url = ? and replaced_by is null and status = 'done' order by id desc", (url,)).fetchone()
    if replaces is not None and not prev:
        raise HTTPException(404, "The source to replace doesn't exist.")
    lineage = prev and s.db.execute("select coalesce(lineage, id) from sources where id = ?", (prev["id"],)).fetchone()[0]
    with WRITE:
        cur = s.db.execute("insert into sources (kind, title, text, status, created_at, lineage, url) "
                           "values ('document', ?, '', 'processing', ?, ?, ?)", (url, now, lineage, url))
    background.add_task(ingest, s.db, s.reader, s.clock, cur.lastrowid, url, url=url)
    return {"id": cur.lastrowid, "status": "processing"}


def _kept(state, source_id, filename):
    folder = state.db_path.parent / "uploads"
    folder.mkdir(exist_ok=True)
    return folder / f"{source_id}{Path(filename).suffix.lower()}"


@router.post("/sources/{id}/retry", status_code=202)
def retry(id: int, background: BackgroundTasks, request: Request):
    """Read a failed upload again, from the copy kept when it was uploaded."""
    s = request.app.state
    src = _source(s.db, id)
    if src["status"] != "failed":
        raise HTTPException(409, "Only a failed upload can be retried.")
    if src["url"]:  # a website: fetch it again
        with WRITE:
            s.db.execute("update sources set status = 'processing', error = null where id = ?", (id,))
        background.add_task(ingest, s.db, s.reader, s.clock, id, src["url"], url=src["url"])
        return {"id": id, "status": "processing"}
    kept = _kept(s, id, src["title"])
    if not kept.exists():
        raise HTTPException(409, "The file wasn't kept (it was uploaded before retrying existed). Upload it again.")
    with WRITE:
        s.db.execute("update sources set status = 'processing', error = null where id = ?", (id,))
    background.add_task(ingest, s.db, s.reader, s.clock, id, src["title"], kept.read_bytes())
    return {"id": id, "status": "processing"}


def _source(con, id):
    r = con.execute("select id, kind, title, status, error, dropped, created_at, lineage, url from sources where id = ?", (id,)).fetchone()
    if not r:
        raise HTTPException(404)
    # what was left out is listed until the student settles the question about it
    # (answered, or "not relevant"): dropped is dropped everywhere
    settled = " ".join(q["text"] for q in con.execute(
        "select text from questions where source_id = ? and status != 'open' and text like 'I couldn''t place%'", (id,)))
    # a drafted question whose quote didn't check out isn't something to place, and isn't asked (ask_left_out)
    return {**dict(r), "dropped": [d for d in json.loads(r["dropped"] or "[]")
                                   if f"“{d['title']}”" not in settled and not d["title"].rstrip().endswith("?")]}


@router.get("/sources")
def list_sources(request: Request):
    con = request.app.state.db
    return [_source(con, r["id"]) for r in con.execute(
        "select id from sources where kind = 'document' and replaced_by is null order by id desc")]


@router.get("/sources/{id}")
def get_source(id: int, request: Request):
    return _source(request.app.state.db, id)
